"""Evaluate every formula cell in the workbook with a small evaluator and compare with independently computed values.

Supported: numbers, plain same-sheet cell references (no absolute markers), ranges inside SUM, SUM(...) with several arguments, and + / -.
Anything else raises FormulaError; an unsupported formula is reported as an audit failure, never skipped.
"""
from __future__ import annotations

import re
from dataclasses import dataclass

import openpyxl

TOLERANCE = 0.01
TOKEN_RE = re.compile(r"\s*(?:(?P<num>\d+(?:\.\d+)?)|(?P<range>[A-Z]{1,3}\d+:[A-Z]{1,3}\d+)"
                      r"|(?P<fn>[A-Za-z]+)\(|(?P<ref>[A-Z]{1,3}\d+)|(?P<op>[-+(),]))")


class FormulaError(Exception):
    pass


def _tokenize(text: str) -> list[tuple[str, str]]:
    tokens, pos = [], 0
    while pos < len(text):
        if text[pos:].strip() == "":
            break
        m = TOKEN_RE.match(text, pos)
        if not m:
            raise FormulaError(f"unsupported syntax near {text[pos:]!r}")
        kind = m.lastgroup
        tokens.append((kind, m.group(kind)))
        pos = m.end()
    return tokens


def evaluate_formula(formula: str, cell, cell_range) -> float:
    """cell(ref) -> number; cell_range(first_ref, last_ref) -> list of numbers."""
    if not formula.startswith("="):
        raise FormulaError("not a formula")
    tokens = _tokenize(formula[1:])
    pos = 0

    def peek():
        return tokens[pos] if pos < len(tokens) else (None, None)

    def take():
        nonlocal pos
        tok = peek()
        pos += 1
        return tok

    def expr() -> float:
        total = term()
        while peek() in (("op", "+"), ("op", "-")):
            op = take()[1]
            rhs = term()
            total = total + rhs if op == "+" else total - rhs
        return total

    def term() -> float:
        kind, value = take()
        if kind == "op" and value == "-":
            return -term()
        if kind == "op" and value == "+":
            return term()
        if kind == "num":
            return float(value)
        if kind == "ref":
            return cell(value)
        if kind == "op" and value == "(":
            inner = expr()
            if take() != ("op", ")"):
                raise FormulaError("missing closing parenthesis")
            return inner
        if kind == "fn":
            if value.upper() != "SUM":
                raise FormulaError(f"unsupported function {value}")
            return sum_args()
        raise FormulaError(f"unexpected token {value!r}")

    def sum_args() -> float:
        total = 0.0
        while True:
            kind, value = peek()
            if kind == "range":
                take()
                first, last = value.split(":")
                total += sum(cell_range(first, last))
            else:
                total += expr()
            sep = take()
            if sep == ("op", ")"):
                return total
            if sep != ("op", ","):
                raise FormulaError("malformed SUM arguments")

    result = expr()
    if pos != len(tokens):
        raise FormulaError("unsupported syntax: trailing tokens")
    return result


@dataclass(frozen=True)
class AuditResult:
    checked: int
    errors: tuple[str, ...]

    @property
    def ok(self) -> bool:
        return not self.errors


def _split(ref: str) -> tuple[int, int]:
    m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
    col = 0
    for ch in m.group(1):
        col = col * 26 + ord(ch) - 64
    return col, int(m.group(2))


def audit_workbook(path, expected: dict[tuple[str, str], float]) -> AuditResult:
    """expected maps (sheet, cell) to the value computed outside the spreadsheet for every formula cell."""
    wb = openpyxl.load_workbook(path)
    errors: list[str] = []
    checked = 0
    formula_cells: set[tuple[str, str]] = set()
    for ws in wb.worksheets:
        grid = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
        memo: dict[str, float] = {}
        active: set[str] = set()

        def value_of(ref: str, grid=grid, memo=memo, active=active, sheet=ws.title) -> float:
            if ref in memo:
                return memo[ref]
            raw = grid.get(ref)
            if raw is None:
                return 0.0
            if isinstance(raw, str) and raw.startswith("="):
                if ref in active:
                    raise FormulaError(f"circular reference at {ref}")
                active.add(ref)
                memo[ref] = evaluate_formula(raw, value_of, lambda a, b: range_values(a, b))
                active.discard(ref)
                return memo[ref]
            if isinstance(raw, (int, float)) and not isinstance(raw, bool):
                return float(raw)
            raise FormulaError(f"{sheet}!{ref} holds text {raw!r} but is used as a number")

        def range_values(first: str, last: str, grid=grid, value_of=value_of) -> list[float]:
            (c1, r1), (c2, r2) = _split(first), _split(last)
            out = []
            for r in range(min(r1, r2), max(r1, r2) + 1):
                for c in range(min(c1, c2), max(c1, c2) + 1):
                    coord = f"{openpyxl.utils.get_column_letter(c)}{r}"
                    raw = grid.get(coord)
                    if isinstance(raw, str) and not raw.startswith("="):
                        continue  # text in a SUM range is ignored, as in a spreadsheet
                    out.append(value_of(coord))
            return out

        for ref, raw in grid.items():
            if not (isinstance(raw, str) and raw.startswith("=")):
                continue
            formula_cells.add((ws.title, ref))
            if (ws.title, ref) not in expected:
                errors.append(f"{ws.title}!{ref}: formula {raw} has no expected value")
                continue
            try:
                got = value_of(ref)
            except FormulaError as exc:
                errors.append(f"{ws.title}!{ref}: unsupported or invalid formula {raw}: {exc}")
                continue
            checked += 1
            if abs(got - expected[(ws.title, ref)]) > TOLERANCE:
                errors.append(f"{ws.title}!{ref}: formula {raw} evaluates to {got:.2f}, expected "
                              f"{expected[(ws.title, ref)]:.2f}")
    for key in expected:
        if key not in formula_cells:
            errors.append(f"{key[0]}!{key[1]}: an expected value was given but the cell is not a formula")
    return AuditResult(checked, tuple(errors))
