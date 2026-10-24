"""The Excel management pack, built from the mart tables only.

Sheets: a P&L by month for each fiscal year (live SUM formulas for section totals, gross margin and operating income),
a store P&L, budget against actual, run rate and variance, and AP (reconciliation and aging). build_pack returns the
expected value of every formula cell, computed in Python from the same rows, so audit_workbook can check the
spreadsheet's arithmetic against an independent figure.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as col

MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDE6EE")
NUM = '#,##0.00;-#,##0.00;"-"'
PCT = "0.0%"
FIRST_MONTH_COL = 4  # D
NOTE = "All data in this pack is synthetic. Amounts are NIU (normalized index units)."

Expected = dict[tuple[str, str], float]


def _rows(con, sql: str, params=None) -> list[dict]:
    cur = con.execute(sql, params or [])
    names = [d[0] for d in cur.description]
    return [{k: (float(v) if hasattr(v, "as_tuple") else v) for k, v in zip(names, r, strict=True)} for r in cur.fetchall()]


def _head(ws, row: int, labels) -> None:
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill = BOLD, HEAD_FILL


def _title(ws, text: str, extra: str | None = None) -> None:
    ws["A1"] = f"Larkspur Outfitters (fictional) - {text}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = NOTE
    if extra:
        ws["A3"] = extra


def _formula(ws, ref: str, text: str, value: float, exp: Expected) -> None:
    ws[ref] = text
    ws[ref].number_format = NUM
    exp[(ws.title, ref)] = round(value, 2)


def _pl_sheet(wb, year: int, rows: list[dict], close: date, exp: Expected) -> None:
    ws = wb.create_sheet(f"P&L FY{year}")
    partial = {r["month_start"].month for r in rows if r["is_partial_month"]}
    _title(ws, f"management P&L FY{year}", "Income and expense lines are shown as positive amounts. Months after the "
           "close month are blank. A partial month appears here but is not in any run rate." if partial else
           "Income and expense lines are shown as positive amounts. Months after the close month are blank.")
    labels = [f"{MONTHS[m - 1]} (partial)" if m in partial else MONTHS[m - 1] for m in range(1, 13)]
    _head(ws, 4, ["Reporting line", "Section", "Type"] + labels + ["Year total"])
    last = FIRST_MONTH_COL + 11
    total_col = last + 1
    by_section: dict[tuple[int, str], list[dict]] = {}
    for r in rows:
        by_section.setdefault((r["section_order"], r["section_label"]), []).append(r)
    values: dict[tuple[str, int], float] = {}  # (line name, month) -> amount
    row_no = 5
    totals: list[tuple[str, str, int, dict[int, float]]] = []  # (label, sign, row, month values)
    for (order, label), section_rows in sorted(by_section.items()):
        first = row_no
        sign = section_rows[0]["sign"]
        for name in sorted({r["line_name"] for r in section_rows}, key=lambda n: min(
                r["line_order"] for r in section_rows if r["line_name"] == n)):
            ws.cell(row=row_no, column=1, value=name)
            ws.cell(row=row_no, column=2, value=label)
            ws.cell(row=row_no, column=3, value=sign)
            line_total = 0.0
            for r in section_rows:
                if r["line_name"] != name:
                    continue
                m = r["month_start"].month
                c = ws.cell(row=row_no, column=FIRST_MONTH_COL + m - 1, value=round(r["amount"], 2))
                c.number_format = NUM
                values[(name, m)] = round(r["amount"], 2)
                line_total += round(r["amount"], 2)
            _formula(ws, f"{col(total_col)}{row_no}", f"=SUM({col(FIRST_MONTH_COL)}{row_no}:{col(last)}{row_no})",
                     line_total, exp)
            row_no += 1
        month_vals: dict[int, float] = {}
        ws.cell(row=row_no, column=1, value=f"Total {label.lower()}").font = BOLD
        for m in range(1, 13):
            month_vals[m] = sum(v for (n, mm), v in values.items() if mm == m and n in
                                {r["line_name"] for r in section_rows})
            letter = col(FIRST_MONTH_COL + m - 1)
            _formula(ws, f"{letter}{row_no}", f"=SUM({letter}{first}:{letter}{row_no - 1})", month_vals[m], exp)
        _formula(ws, f"{col(total_col)}{row_no}", f"=SUM({col(FIRST_MONTH_COL)}{row_no}:{col(last)}{row_no})",
                 sum(month_vals.values()), exp)
        totals.append((label, sign, row_no, month_vals))
        row_no += 2
    income = [t for t in totals if t[1] == "income"]
    expense = [t for t in totals if t[1] == "expense"]
    if income and expense:
        def combine(label: str, plus: list, minus: list) -> None:
            nonlocal row_no
            ws.cell(row=row_no, column=1, value=label).font = BOLD
            vals = {}
            for m in [*range(1, 13), 0]:  # months first, then the year-total column
                if m == 0:
                    cidx, key = total_col, None
                else:
                    cidx, key = FIRST_MONTH_COL + m - 1, m
                letter = col(cidx)
                text = "=" + "+".join(f"{letter}{p[2]}" for p in plus) + "".join(f"-{letter}{q[2]}" for q in minus)
                if key is None:
                    value = sum(vals.values())
                else:
                    value = sum(p[3][key] for p in plus) - sum(q[3][key] for q in minus)
                    vals[key] = value
                _formula(ws, f"{letter}{row_no}", text, value, exp)
            row_no += 1
            return vals

        gm_vals = combine("Gross margin", income, expense[:1])
        gm_row = row_no - 1
        # operating income = gross margin less the remaining expense sections
        ws.cell(row=row_no, column=1, value="Operating income").font = BOLD
        for m in [*range(1, 13), 0]:
            cidx = total_col if m == 0 else FIRST_MONTH_COL + m - 1
            letter = col(cidx)
            text = f"={letter}{gm_row}" + "".join(f"-{letter}{q[2]}" for q in expense[1:])
            if m == 0:
                value = sum(gm_vals[k] - sum(q[3][k] for q in expense[1:]) for k in range(1, 13))
            else:
                value = gm_vals[m] - sum(q[3][m] for q in expense[1:])
            _formula(ws, f"{letter}{row_no}", text, value, exp)
        row_no += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    ws.freeze_panes = "D5"


def _store_sheet(wb, con, year: int, close: date, exp: Expected) -> None:
    ws = wb.create_sheet("Store P&L")
    _title(ws, f"store P&L, FY{year} year to date (closed months to {MONTHS[close.month - 1]})",
           "Head Office and Online carries online sales and head office expense.")
    _head(ws, 4, ["Cost centre", "Name", "Revenue", "Cost of goods", "Gross margin", "Gross margin %", "Store operating expense",
                  "Head office expense", "Operating income"])
    data = _rows(con, """
        select cost_center_code, any_value(cost_center_name) as name, sum(revenue) as revenue, sum(cogs) as cogs,
               sum(store_opex) as store_opex, sum(ho_opex) as ho_opex
        from marts.rpt_store_pl where fiscal_year = ? and month_start <= ? and not is_partial_month
        group by cost_center_code order by cost_center_code""", [year, close])
    row_no = 5
    first = row_no
    sums = {"C": 0.0, "D": 0.0, "E": 0.0, "G": 0.0, "H": 0.0, "I": 0.0}
    for d in data:
        rev, cogs = round(d["revenue"] or 0, 2), round(d["cogs"] or 0, 2)
        so, ho = round(d["store_opex"], 2), round(d["ho_opex"], 2)
        ws.cell(row=row_no, column=1, value=d["cost_center_code"])
        ws.cell(row=row_no, column=2, value=d["name"])
        for letter, v in (("C", rev), ("D", cogs), ("G", so), ("H", ho)):
            ws[f"{letter}{row_no}"] = v
            ws[f"{letter}{row_no}"].number_format = NUM
            sums[letter] += v
        _formula(ws, f"E{row_no}", f"=C{row_no}-D{row_no}", rev - cogs, exp)
        ws[f"F{row_no}"] = (rev - cogs) / rev if rev else None
        ws[f"F{row_no}"].number_format = PCT
        _formula(ws, f"I{row_no}", f"=E{row_no}-G{row_no}-H{row_no}", rev - cogs - so - ho, exp)
        sums["E"] += rev - cogs
        sums["I"] += rev - cogs - so - ho
        row_no += 1
    ws.cell(row=row_no, column=1, value="Total").font = BOLD
    for letter in "CDEGHI":
        _formula(ws, f"{letter}{row_no}", f"=SUM({letter}{first}:{letter}{row_no - 1})" if row_no > first else "=0",
                 sums[letter], exp)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 24


def _budget_sheet(wb, con, year: int, close: date, exp: Expected) -> None:
    ws = wb.create_sheet("Budget vs Actual")
    _title(ws, f"budget against actual, FY{year} year to date (closed months to {MONTHS[close.month - 1]})",
           "Variance is actual less budget. A miss is adverse by more than 10 percent of budget and more than "
           "3,000 NIU in a quarter.")
    _head(ws, 4, ["Section", "Budget", "Actual", "Variance (actual less budget)", "Variance %"])
    data = _rows(con, """
        select section_label, sum(budget_amount) as budget, sum(actual_amount) as actual
        from marts.rpt_budget_vs_actual where fiscal_year = ? and month_start <= ? and not is_partial_month
        group by section_label, section_key order by min(case section_key when 'revenue' then 1 when 'cogs' then 2
            when 'store_opex' then 3 else 4 end)""", [year, close])
    row_no = 5
    for d in data:
        budget, actual = round(d["budget"], 2), round(d["actual"], 2)
        ws.cell(row=row_no, column=1, value=d["section_label"])
        ws[f"B{row_no}"], ws[f"C{row_no}"] = budget, actual
        ws[f"B{row_no}"].number_format = ws[f"C{row_no}"].number_format = NUM
        _formula(ws, f"D{row_no}", f"=C{row_no}-B{row_no}", actual - budget, exp)
        ws[f"E{row_no}"] = (actual - budget) / budget if budget else None
        ws[f"E{row_no}"].number_format = PCT
        row_no += 1
    row_no += 1
    ws.cell(row=row_no, column=1, value="Quarterly budget misses").font = BOLD
    row_no += 1
    _head(ws, row_no, ["Quarter starting", "Line", "Cost centre", "Budget", "Actual", "Adverse amount", "Adverse %"])
    for d in _rows(con, """select quarter_start, line_name, cost_center_code, budget_amount, actual_amount, adverse_amount,
                           adverse_pct from marts.rpt_budget_misses where is_miss order by quarter_start, line_name,
                           cost_center_code"""):
        row_no += 1
        vals = [d["quarter_start"], d["line_name"], d["cost_center_code"], d["budget_amount"], d["actual_amount"],
                d["adverse_amount"], d["adverse_pct"]]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=row_no, column=i, value=v)
            if i in (4, 5, 6):
                c.number_format = NUM
            if i == 7:
                c.number_format = PCT
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 32


def _run_rate_sheet(wb, con, exp: Expected) -> None:
    ws = wb.create_sheet("Run rate & variance")
    window = _rows(con, "select min(window_start) as s, max(window_end) as e, max(months_in_window) as n from marts.rpt_run_rate")
    note = ""
    if window and window[0]["s"]:
        note = (f"Run rate is the average month over the {window[0]['n']} closed months from {window[0]['s']:%b %Y} to "
                f"{window[0]['e']:%b %Y}. Journals on the excluded list stay in the P&L and come out here.")
    _title(ws, "run rate and variance", note)
    _head(ws, 4, ["Reporting line", "Section", "Run rate per month", "Run rate incl. excluded journals",
                  "Effect of excluded journals", "Current period per month", "Prior period per month",
                  "Change per month", "Change %", "Flagged", "Change % incl. excluded", "Flagged incl. excluded"])
    var = {v["line_name"]: v for v in _rows(con, "select * from marts.rpt_variance")}
    row_no = 5
    for r in _rows(con, "select * from marts.rpt_run_rate order by section_order, line_order"):
        v = var.get(r["line_name"])
        ws.cell(row=row_no, column=1, value=r["line_name"])
        ws.cell(row=row_no, column=2, value=r["section_label"])
        ws[f"C{row_no}"], ws[f"D{row_no}"] = round(r["monthly_run_rate"], 2), round(r["monthly_run_rate_incl_excluded"], 2)
        _formula(ws, f"E{row_no}", f"=D{row_no}-C{row_no}",
                 round(r["monthly_run_rate_incl_excluded"], 2) - round(r["monthly_run_rate"], 2), exp)
        if v:
            cur, pri = round(v["current_monthly"], 2), round(v["prior_monthly"], 2) if v["prior_monthly"] is not None else None
            ws[f"F{row_no}"] = cur
            ws[f"G{row_no}"] = pri
            if pri is not None:
                _formula(ws, f"H{row_no}", f"=F{row_no}-G{row_no}", cur - pri, exp)
            ws[f"I{row_no}"] = v["change_pct"]
            ws[f"J{row_no}"] = "yes" if v["is_flagged"] else ""
            ws[f"K{row_no}"] = v["change_pct_gross"]
            ws[f"L{row_no}"] = "yes" if v["is_flagged_gross"] else ""
        for letter in "CDFG":
            ws[f"{letter}{row_no}"].number_format = NUM
        for letter in "IK":
            ws[f"{letter}{row_no}"].number_format = PCT
        row_no += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24


def _ap_sheet(wb, con, exp: Expected) -> None:
    ws = wb.create_sheet("AP")
    _title(ws, "accounts payable: reconciliation, aging and days payable outstanding",
           "The bill-pay platform and the accounting system both record most bills. The naive total counts them twice.")
    _head(ws, 4, ["Reconciliation", "Amount"])
    t = _rows(con, "select * from marts.rpt_ap_totals")
    t = t[0] if t else {"platform_total": 0.0, "accounting_total": 0.0, "matched_total": 0.0, "held_out_total": 0.0,
                        "deduplicated_total": 0.0, "matched_count": 0, "platform_only_count": 0,
                        "accounting_only_count": 0, "ambiguous_count": 0}
    ws["A5"], ws["B5"] = "Platform bills total", round(t["platform_total"], 2)
    ws["A6"], ws["B6"] = "Accounting bill lines total", round(t["accounting_total"], 2)
    ws["A7"] = "Naive total"
    _formula(ws, "B7", "=B5+B6", round(t["platform_total"], 2) + round(t["accounting_total"], 2), exp)
    ws["A8"], ws["B8"] = "Less platform bills matched to a document", round(t["matched_total"], 2)
    ws["A9"], ws["B9"] = "Less ambiguous platform bills held out", round(t["held_out_total"], 2)
    ws["A10"] = "Deduplicated total"
    ws["A10"].font = BOLD
    naive = round(t["platform_total"], 2) + round(t["accounting_total"], 2)
    _formula(ws, "B10", "=B7-B8-B9", naive - round(t["matched_total"], 2) - round(t["held_out_total"], 2), exp)
    for ref in ("B5", "B6", "B8", "B9"):
        ws[ref].number_format = NUM
    _head(ws, 12, ["Bills by outcome", "Count"])
    for i, (label, key) in enumerate((("Matched", "matched_count"), ("Platform only (never synced)", "platform_only_count"),
                                      ("Accounting only (entered directly)", "accounting_only_count"),
                                      ("Ambiguous (held out)", "ambiguous_count")), start=13):
        ws[f"A{i}"], ws[f"B{i}"] = label, t[key]
    latest = _rows(con, "select max(as_of) as d from marts.rpt_ap_aging")
    row_no = 19
    if latest and latest[0]["d"]:
        as_of = latest[0]["d"]
        _head(ws, row_no, [f"Open payables at {as_of:%d %b %Y}", "Bills", "Amount"])
        first = row_no + 1
        total = 0.0
        for d in _rows(con, "select bucket, bill_count, open_amount from marts.rpt_ap_aging where as_of = ? order by bucket_order", [as_of]):
            row_no += 1
            ws[f"A{row_no}"], ws[f"B{row_no}"], ws[f"C{row_no}"] = d["bucket"], d["bill_count"], round(d["open_amount"], 2)
            ws[f"C{row_no}"].number_format = NUM
            total += round(d["open_amount"], 2)
        row_no += 1
        ws[f"A{row_no}"] = "Total open"
        ws[f"A{row_no}"].font = BOLD
        _formula(ws, f"C{row_no}", f"=SUM(C{first}:C{row_no - 1})", total, exp)
        row_no += 2
    _head(ws, row_no, ["Month", "Open payables", "Bills in last 90 days", "DPO (days)", "Days to pay (bills paid in month)",
                       "Share paid late"])
    for d in _rows(con, "select * from (select * from marts.rpt_dpo order by month_start desc limit 12) order by month_start"):
        row_no += 1
        vals = [d["month_start"].strftime("%Y-%m"), round(d["open_amount"], 2), round(d["bills_90d"], 2), d["dpo"],
                d["avg_days_to_pay"], d["share_paid_late"]]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=row_no, column=i, value=v)
            if i in (2, 3):
                c.number_format = NUM
            elif i in (4, 5):
                c.number_format = "0.0"
            elif i == 6:
                c.number_format = PCT
    ws.column_dimensions["A"].width = 42
    for letter in "BCDEF":
        ws.column_dimensions[letter].width = 22


def build_pack(con, path, month: date) -> Expected:
    """Write the pack for a close month and return the expected value of every formula cell."""
    exp: Expected = {}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    pl = _rows(con, """select month_start, fiscal_year, is_partial_month, section_label, section_order, line_name,
                       line_order, sign, amount from marts.rpt_monthly_pl where month_start <= ?
                       order by fiscal_year, section_order, line_order, month_start""", [month])
    years = sorted({r["fiscal_year"] for r in pl})
    for year in years:
        _pl_sheet(wb, year, [r for r in pl if r["fiscal_year"] == year], month, exp)
    _store_sheet(wb, con, month.year, month, exp)
    _budget_sheet(wb, con, month.year, month, exp)
    _run_rate_sheet(wb, con, exp)
    _ap_sheet(wb, con, exp)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return exp
