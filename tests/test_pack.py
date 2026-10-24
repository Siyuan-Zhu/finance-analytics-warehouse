from datetime import date

import duckdb
import openpyxl
import pytest

from src.audit import audit_workbook
from src.pack import build_pack
from tests.marts_fixture import FEB, JAN, MAR, line, make_marts


def monthly_rows():
    rows = []
    for month, apparel, online, product, wages in ((JAN, 1000, 400, 600, 300), (FEB, 1100, 500, 650, 310), (MAR, 900, 450, 500, 305)):
        partial = month == MAR
        rows += [
            line(month, "rev", "Apparel Sales", "Revenue", 1, "income", apparel, partial, line_order=1, section_order=1),
            line(month, "rev", "Online Sales", "Revenue", 1, "income", online, partial, line_order=2, section_order=1),
            line(month, "cogs", "Product Cost", "Cost of goods", 2, "expense", product, partial, line_order=1, section_order=2),
            line(month, "store", "Store Wages", "Store operating expense", 3, "expense", wages, partial, line_order=1, section_order=3),
        ]
    return rows


@pytest.fixture
def con():
    c = duckdb.connect(":memory:")
    make_marts(c, rpt_monthly_pl=monthly_rows())
    return c


def cell(path, sheet, ref):
    return openpyxl.load_workbook(path)[sheet][ref].value


def test_pack_has_the_expected_sheets(con, tmp_path):
    path = tmp_path / "pack.xlsx"
    build_pack(con, path, date(2025, 2, 1))
    assert openpyxl.load_workbook(path).sheetnames == [
        "P&L FY2025", "Store P&L", "Budget vs Actual", "Run rate & variance", "AP"]


def test_section_totals_are_live_sum_formulas(con, tmp_path):
    path = tmp_path / "pack.xlsx"
    build_pack(con, path, date(2025, 2, 1))
    ws = openpyxl.load_workbook(path)["P&L FY2025"]
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=SUM(")]
    assert len(formulas) >= 3 * 12  # at least one section-total formula per month


def test_every_formula_matches_an_independent_expected_value(con, tmp_path):
    path = tmp_path / "pack.xlsx"
    expected = build_pack(con, path, date(2025, 2, 1))
    result = audit_workbook(path, expected)
    assert result.ok, result.errors
    assert result.checked == len(expected) > 40


def test_known_totals_land_in_the_right_cells(con, tmp_path):
    path = tmp_path / "pack.xlsx"
    expected = build_pack(con, path, date(2025, 2, 1))
    # Operating income for Jan: 1000 + 400 - 600 - 300 = 500. Find the row labelled by the pack, not by position.
    ws = openpyxl.load_workbook(path)["P&L FY2025"]
    row = next(r for r in ws.iter_rows() if r[0].value == "Operating income")
    assert expected[("P&L FY2025", f"D{row[0].row}")] == pytest.approx(500.0)
    assert expected[("P&L FY2025", f"E{row[0].row}")] == pytest.approx(1100 + 500 - 650 - 310)


def test_months_after_the_close_month_are_left_empty_and_partial_is_labelled(con, tmp_path):
    path = tmp_path / "pack.xlsx"
    build_pack(con, path, date(2025, 3, 1))
    ws = openpyxl.load_workbook(path)["P&L FY2025"]
    header = next(r for r in ws.iter_rows() if r[0].value == "Reporting line")
    assert header[5].value == "Mar (partial)"
    assert header[6].value == "Apr"
    line_row = next(r for r in ws.iter_rows() if r[0].value == "Apparel Sales")
    assert line_row[6].value is None


def test_pack_says_the_data_is_synthetic_and_uses_niu(con, tmp_path):
    path = tmp_path / "pack.xlsx"
    build_pack(con, path, date(2025, 2, 1))
    text = " ".join(str(c.value) for ws in openpyxl.load_workbook(path).worksheets for row in ws.iter_rows() for c in row if c.value)
    assert "synthetic" in text.lower() and "NIU" in text
    assert not any(sym in text for sym in "$£€¥")


def test_ap_sheet_dedup_total_is_a_formula(tmp_path):
    c = duckdb.connect(":memory:")
    make_marts(c, rpt_monthly_pl=monthly_rows(),
               rpt_ap_totals=[(1000.0, 800.0, 600.0, 50.0, 1150.0, 6, 2, 3, 1)])
    path = tmp_path / "pack.xlsx"
    expected = build_pack(c, path, date(2025, 2, 1))
    assert audit_workbook(path, expected).ok
    ws = openpyxl.load_workbook(path)["AP"]
    row = next(r for r in ws.iter_rows() if r[0].value == "Deduplicated total")
    assert str(row[1].value).startswith("=")
    assert expected[("AP", f"B{row[0].row}")] == pytest.approx(1150.0)
