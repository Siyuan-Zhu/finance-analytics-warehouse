import dataclasses

import openpyxl
import pytest

from src.gl_parser import parse_gl
from tests.conftest import txn, write_gl


@pytest.fixture
def small_gl(tmp_path):
    return write_gl(
        tmp_path / "gl.xlsx",
        [
            ("4045", "Apparel Sales", [txn(1, "S-1", "Weekly sales", "ST01", credit=500.0)]),
            ("6105", "Store Wages", [txn(1, "P-1", "Payroll", "ST01", debit=120.5), txn(2, "P-2", "Payroll", "ST02", debit=99.5)]),
            ("1010", "Operating Bank Account", [txn(1, "S-1", "Receipt", "ST01", debit=500.0)]),
            ("1900", "Inter-account Transfers", [txn(1, "T-1", "Sweep to savings", "HQ01", debit=50.0)]),
        ],
    )


def test_rows_are_typed_and_signed(small_gl):
    parsed = parse_gl(small_gl)
    by_ref = {t.reference: t for t in parsed.transactions}
    wages = [t for t in parsed.transactions if t.account_code == "6105"]
    assert sum(t.amount for t in wages) == pytest.approx(220.0)
    sales = next(t for t in parsed.transactions if t.account_code == "4045")
    assert sales.amount == -500.0  # credit is negative: amount = debit - credit
    assert sales.account_name == "Apparel Sales"
    assert sales.cost_center == "ST01"
    assert by_ref["T-1"].account_code == "1900"
    assert sales.entry_date.isoformat() == "2025-01-10"


def test_disposition_counts(small_gl):
    parsed = parse_gl(small_gl)
    assert parsed.dispositions == {
        "title": 3,
        "blank": 4,
        "account_header": 4,
        "banner": 4,
        "subtotal": 4,
        "transaction": 5,
    }
    assert parsed.errors == ()


def test_records_are_immutable(small_gl):
    parsed = parse_gl(small_gl)
    with pytest.raises(dataclasses.FrozenInstanceError):
        parsed.transactions[0].amount = 1  # type: ignore[misc]


def test_subtotal_mismatch_is_reported(tmp_path):
    path = write_gl(tmp_path / "gl.xlsx", [("6105", "Store Wages", [txn(1, "P-1", "Payroll", "ST01", debit=100.0)])])
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for row in ws.iter_rows():
        if row[0].value == "Total 6105 Store Wages":
            row[4].value = 999.0
    wb.save(path)
    parsed = parse_gl(path)
    assert any("subtotal" in e.lower() and "6105" in e for e in parsed.errors)


def test_unrecognised_row_is_an_error_not_a_skip(tmp_path):
    path = write_gl(tmp_path / "gl.xlsx", [("6105", "Store Wages", [txn(1, "P-1", "Payroll", "ST01", debit=100.0)])])
    wb = openpyxl.load_workbook(path)
    wb.active.append(["Some stray footer text that is not a row type"])
    wb.save(path)
    parsed = parse_gl(path)
    assert parsed.dispositions["unrecognised"] == 1
    assert len(parsed.errors) == 1


def test_transaction_before_any_account_header_is_an_error(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Title"])
    ws.append([txn(1, "X-1", "orphan", "ST01", debit=1.0)[0], "X-1", "orphan", "ST01", 1.0, None])
    wb.save(tmp_path / "gl.xlsx")
    parsed = parse_gl(tmp_path / "gl.xlsx")
    assert parsed.transactions == ()
    assert any("before" in e.lower() for e in parsed.errors)


def test_account_header_while_previous_account_is_open_is_an_error(tmp_path):
    path = write_gl(tmp_path / "gl.xlsx", [
        ("6105", "Store Wages", [txn(1, "P-1", "Payroll", "ST01", debit=100.0)]),
        ("6145", "Store Rent", [txn(1, "R-1", "Rent", "ST01", debit=50.0)]),
    ])
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for row in ws.iter_rows():
        if row[0].value == "Total 6105 Store Wages":
            for c in row:
                c.value = None  # the first account loses its subtotal row
    wb.save(path)
    parsed = parse_gl(path)
    assert any("6105" in e and "subtotal" in e.lower() for e in parsed.errors)


def test_account_without_subtotal_at_end_of_file_is_an_error(tmp_path):
    path = write_gl(tmp_path / "gl.xlsx", [("6105", "Store Wages", [txn(1, "P-1", "Payroll", "ST01", debit=100.0)])])
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for row in ws.iter_rows():
        if row[0].value == "Total 6105 Store Wages":
            for c in row:
                c.value = None
    wb.save(path)
    parsed = parse_gl(path)
    assert any("6105" in e and "subtotal" in e.lower() for e in parsed.errors)
